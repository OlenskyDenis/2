"""Наскрізні тести валідації сценаріїв з quickstart.md (T026)."""
from datetime import date
from pathlib import Path
import openpyxl
import pytest

from src.bond_collector.excel import ExcelFileLockedError, ExcelStorage
from src.bond_collector.models import Bond, CashFlowEvent, CashFlowType
from src.bond_collector.service import CollectorService
from src.cli.main import main


@pytest.mark.integration
class TestQuickstartScenarios:
    """Перевірка сценаріїв валідації, описаних у quickstart.md."""

    def test_quickstart_auto_creation_and_three_sheets(self, tmp_path: Path, monkeypatch):
        """Сценарій 1: Відсутність файлу Excel при першому запуску -> створення з нуля з 3 аркушами."""
        target_file = tmp_path / "bonds.xlsx"
        assert not target_file.exists()

        # Мокуємо отримання тестових облігацій
        test_bond = Bond(
            isin="UA4000227185",
            name="Військові облігації",
            currency="UAH",
            maturity_date=date(2027, 5, 26),
            is_military=True,
            price_buy=1056.23,
            rate_buy=14.85,
            price_sell=995.10,
            rate_sell=17.50,
            available_qty=85217,
            min_amount=1,
            max_amount=85217,
            cash_flows=[
                CashFlowEvent(
                    isin="UA4000227185",
                    payment_date=date(2025, 5, 21),
                    event_type=CashFlowType.COUPON,
                    amount=75.50,
                    currency="UAH",
                )
            ],
        )

        monkeypatch.setattr(CollectorService, "collect_all", lambda self, progress_callback=None: [test_bond])

        ret = main(["--output", str(target_file)])
        assert ret == 0
        assert target_file.exists()

        wb = openpyxl.load_workbook(target_file)
        assert len(wb.sheetnames) == 3
        assert "Актуальні котирування" in wb.sheetnames
        assert "Історія котирувань" in wb.sheetnames
        assert "Графік виплат" in wb.sheetnames

        # Перевірка даних на першому аркуші
        ws_curr = wb["Актуальні котирування"]
        assert ws_curr.cell(row=2, column=1).value == "UA4000227185"
        assert ws_curr.cell(row=2, column=11).value == 85217
        assert ws_curr.cell(row=2, column=12).value == "В наявності"
        assert ws_curr.auto_filter.ref is not None

    def test_quickstart_file_locked_handling(self, tmp_path: Path, monkeypatch, capsys):
        """Сценарій 2: Файл Excel заблоковано користувачем -> чисте повідомлення у stderr та код 1."""
        target_file = tmp_path / "bonds.xlsx"

        monkeypatch.setattr(CollectorService, "collect_all", lambda self, progress_callback=None: [])

        def mock_locked_save(self, bonds, filepath):
            raise ExcelFileLockedError("Файл 'bonds.xlsx' заблоковано іншою програмою.")

        monkeypatch.setattr(ExcelStorage, "save", mock_locked_save)

        ret = main(["--output", str(target_file)])
        assert ret == 1

        captured = capsys.readouterr()
        assert "[ПОМИЛКА]" in captured.err
        assert "заблоковано" in captured.err
