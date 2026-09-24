"""Модульні тести для ExcelStorage (T014)."""
from datetime import date, datetime
from pathlib import Path
import openpyxl
import pytest

from src.bond_collector.excel import ExcelStorage
from src.bond_collector.models import Bond, CashFlowEvent, CashFlowType


@pytest.mark.unit
class TestExcelStorage:
    """Тестування операцій створення, оновлення, дедуплікації та upsert у Excel."""

    def test_create_new_file_from_scratch(self, tmp_path: Path):
        file_path = tmp_path / "bonds.xlsx"
        storage = ExcelStorage()
        assert not storage.exists(file_path)

        storage.save([], file_path)
        assert storage.exists(file_path)

        wb = openpyxl.load_workbook(file_path)
        assert len(wb.sheetnames) == 3

    def test_save_and_update_bonds(self, tmp_path: Path):
        file_path = tmp_path / "bonds.xlsx"
        storage = ExcelStorage()

        bond1 = Bond(
            isin="UA4000227185",
            name="Військові облігації",
            currency="UAH",
            maturity_date=date(2027, 5, 26),
            is_military=True,
            price_buy=1050.00,
            rate_buy=15.0,
            price_sell=990.00,
            rate_sell=17.0,
            available_qty=1000,
            min_amount=1,
            max_amount=1000,
            is_active=True,
            updated_at=datetime(2026, 9, 24, 10, 0, 0),
            cash_flows=[
                CashFlowEvent(
                    isin="UA4000227185",
                    payment_date=date(2025, 5, 21),
                    event_type=CashFlowType.COUPON,
                    amount=75.50,
                    currency="UAH",
                )
            ],
        )

        # Перше збереження
        storage.save([bond1], file_path)

        wb1 = openpyxl.load_workbook(file_path)
        ws_curr1 = wb1["Актуальні котирування"]
        assert ws_curr1.max_row == 2
        assert ws_curr1.cell(row=2, column=1).value == "UA4000227185"
        assert ws_curr1.cell(row=2, column=7).value == "На вітрині (Купівля)"
        assert ws_curr1.cell(row=2, column=8).value == 1050.00
        assert ws_curr1.cell(row=2, column=12).value == 1000

        ws_hist1 = wb1["Історія котирувань"]
        assert ws_hist1.max_row == 2

        ws_ops1 = wb1["Графік виплат"]
        assert ws_ops1.max_row == 2
        assert ws_ops1.cell(row=2, column=1).value == "UA4000227185"

        # Повторне збереження без змін (перевірка дедуплікації історії)
        storage.save([bond1], file_path)
        wb2 = openpyxl.load_workbook(file_path)
        ws_hist2 = wb2["Історія котирувань"]
        # Рядок не повинен дублюватися!
        assert ws_hist2.max_row == 2

        # Оновлення ціни та залишку в той самий день
        bond1_updated = Bond(
            isin="UA4000227185",
            name="Військові облігації",
            currency="UAH",
            maturity_date=date(2027, 5, 26),
            is_military=True,
            price_buy=1060.00,  # ціна змінилася
            rate_buy=14.5,
            price_sell=990.00,
            rate_sell=17.0,
            available_qty=800,  # залишок зменшився
            min_amount=1,
            max_amount=800,
            is_active=True,
            updated_at=datetime(2026, 9, 24, 14, 0, 0),
        )
        storage.save([bond1_updated], file_path)

        wb3 = openpyxl.load_workbook(file_path)
        ws_curr3 = wb3["Актуальні котирування"]
        assert ws_curr3.max_row == 2
        assert ws_curr3.cell(row=2, column=8).value == 1060.00
        assert ws_curr3.cell(row=2, column=12).value == 800

        # В історії запис за цей день оновився новими значеннями
        ws_hist3 = wb3["Історія котирувань"]
        assert ws_hist3.max_row == 2
        assert ws_hist3.cell(row=2, column=4).value == 1060.00

    def test_nondestructive_retention_of_delisted_bonds_and_past_cashflows(self, tmp_path: Path):
        file_path = tmp_path / "bonds.xlsx"
        storage = ExcelStorage()

        bond_active = Bond(
            isin="UA4000227185",
            name="Активна облігація",
            currency="UAH",
            maturity_date=date(2027, 5, 26),
            is_military=True,
            price_buy=1000.0,
            available_qty=500,
            updated_at=datetime(2026, 9, 24, 10, 0, 0),
            cash_flows=[
                CashFlowEvent(
                    isin="UA4000227185",
                    payment_date=date(2024, 5, 21),  # минула виплата
                    event_type=CashFlowType.COUPON,
                    amount=70.0,
                    currency="UAH",
                )
            ],
        )

        bond_to_delist = Bond(
            isin="UA4000227193",
            name="Облігація під зняття",
            currency="UAH",
            maturity_date=date(2026, 12, 1),
            is_military=True,
            price_buy=950.0,
            available_qty=200,
            updated_at=datetime(2026, 9, 24, 10, 0, 0),
        )

        # Зберігаємо обидві
        storage.save([bond_active, bond_to_delist], file_path)

        # Тепер у новому батчі немає bond_to_delist і немає минулої виплати
        bond_active_only = Bond(
            isin="UA4000227185",
            name="Активна облігація",
            currency="UAH",
            maturity_date=date(2027, 5, 26),
            is_military=True,
            price_buy=1010.0,
            available_qty=450,
            updated_at=datetime(2026, 9, 24, 12, 0, 0),
            cash_flows=[
                CashFlowEvent(
                    isin="UA4000227185",
                    payment_date=date(2027, 5, 26),  # нова майбутня виплата
                    event_type=CashFlowType.REDEMPTION,
                    amount=1000.0,
                    currency="UAH",
                )
            ],
        )
        storage.save([bond_active_only], file_path)

        wb = openpyxl.load_workbook(file_path)
        ws_curr = wb["Актуальні котирування"]
        # Обидва папери залишилися в таблиці!
        assert ws_curr.max_row == 3
        # Знятий папір має статус "Знято з торгів / Погашено"
        delisted_status = None
        for row in range(2, ws_curr.max_row + 1):
            if ws_curr.cell(row=row, column=1).value == "UA4000227193":
                delisted_status = ws_curr.cell(row=row, column=13).value
        assert delisted_status == "Знято з торгів / Погашено"

        # У графіку виплат минула виплата збережена, а нова додана!
        ws_ops = wb["Графік виплат"]
        assert ws_ops.max_row == 3
